"""Turn Markdown into real Word, Excel and PowerPoint files.

WHY MARKDOWN. The work model only has to WRITE - headings, bullets, a
table - and this module builds the file. That split is the safety
property: the model never produces code that runs, and never receives a
path to write to. It hands over text; text is all this module accepts.
It is also far easier for a 4B model to write Markdown correctly than to
emit strict JSON or drive a document library.

    Markdown  -> .docx   headings, paragraphs, bullets, numbered lists,
                         tables, bold/italic, quotes, code blocks
    Markdown  -> .xlsx   each table becomes a sheet (the heading above it
                         names the sheet); "=B2*C2" cells are formulas
    Markdown  -> .pptx   "# " deck title, each "## " a slide, bullets
                         under it (indent = sub-bullet)

Files are ALWAYS new: save_path() never returns a path that exists, so a
document Master already has can never be overwritten by this module.

Needs: python-docx, openpyxl, python-pptx
    py -m pip install python-docx openpyxl python-pptx
"""

import os
import re
import time
from pathlib import Path
from typing import List, Optional, Tuple

EXTENSIONS = {"docx": ".docx", "xlsx": ".xlsx", "pptx": ".pptx"}

# Windows-forbidden characters, and the reserved device names.
_BAD_NAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_RESERVED = {"con", "prn", "aux", "nul",
             *(f"com{i}" for i in range(1, 10)),
             *(f"lpt{i}" for i in range(1, 10))}


class DocumentError(Exception):
    """The file could not be built (missing library, unusable content)."""


# ---------------------------------------------------------------------
# Where files go
# ---------------------------------------------------------------------

def output_dir(override: Optional[str] = None) -> Path:
    """Documents\\GreatSage under the user's profile, created on demand.

    A folder of its own on purpose: everything Great Sage makes is in one
    place, and nothing it writes can land in (or over) Master's own files.
    """
    base = Path(override) if override else Path.home() / "Documents" / "GreatSage"
    base.mkdir(parents=True, exist_ok=True)
    return base


def safe_name(title: str, limit: int = 60) -> str:
    name = _BAD_NAME_CHARS.sub(" ", title or "")
    name = re.sub(r"\s+", " ", name).strip(" .")
    name = name[:limit].strip(" .")
    if not name or name.lower() in _RESERVED:
        name = "documento"
    return name


def save_path(kind: str, title: str, directory: Optional[Path] = None) -> Path:
    """A path in `directory` that does not exist yet. Never overwrites."""
    if kind not in EXTENSIONS:
        raise DocumentError("Tipo de documento desconocido: %r" % kind)
    directory = directory or output_dir()
    stem = safe_name(title)
    candidate = directory / (stem + EXTENSIONS[kind])
    if not candidate.exists():
        return candidate
    stamp = time.strftime("%Y-%m-%d_%H%M")
    candidate = directory / ("%s %s%s" % (stem, stamp, EXTENSIONS[kind]))
    n = 2
    while candidate.exists():
        candidate = directory / ("%s %s (%d)%s"
                                 % (stem, stamp, n, EXTENSIONS[kind]))
        n += 1
    return candidate


# ---------------------------------------------------------------------
# Markdown parsing (shared)
# ---------------------------------------------------------------------

_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+\u2022]\s+(.*)$")
_NUMBERED = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_RULE = re.compile(r"^\s*([-*_])\s*(\1\s*){2,}$")
_INLINE = re.compile(r"(\*\*[^*\n]+\*\*|__[^_\n]+__|\*[^*\n]+\*|`[^`\n]+`)")


def _split_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _clean_inline(text: str) -> str:
    """Text with the Markdown emphasis marks removed (for plain cells)."""
    return _INLINE.sub(lambda m: m.group(0).strip("*_`"), text)


def _blocks(md: str):
    """Yield (type, payload) for each block of the Markdown.

    Types: heading (level, text), bullet (depth, text), numbered
    (depth, text), table (rows), quote (text), code (lines), rule, para.
    """
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i, n = 0, len(lines)
    para: List[str] = []

    def flush():
        if para:
            yield ("para", " ".join(p.strip() for p in para))
            para.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            yield from flush()
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            yield ("code", code)
            continue

        if not stripped:
            yield from flush()
            i += 1
            continue

        if _TABLE_ROW.match(line):
            yield from flush()
            rows = []
            while i < n and _TABLE_ROW.match(lines[i]):
                if not _TABLE_SEP.match(lines[i]):
                    rows.append(_split_row(lines[i]))
                i += 1
            if rows:
                yield ("table", rows)
            continue

        m = _HEADING.match(stripped)
        if m:
            yield from flush()
            yield ("heading", (len(m.group(1)), m.group(2)))
            i += 1
            continue

        if _RULE.match(stripped):
            yield from flush()
            yield ("rule", None)
            i += 1
            continue

        m = _BULLET.match(line)
        if m:
            yield from flush()
            yield ("bullet", (len(m.group(1).expandtabs(2)) // 2, m.group(2)))
            i += 1
            continue

        m = _NUMBERED.match(line)
        if m:
            yield from flush()
            yield ("numbered", (len(m.group(1).expandtabs(2)) // 2, m.group(2)))
            i += 1
            continue

        if stripped.startswith(">"):
            yield from flush()
            yield ("quote", stripped.lstrip("> ").strip())
            i += 1
            continue

        para.append(line)
        i += 1
    yield from flush()


def first_heading(md: str) -> str:
    for kind, payload in _blocks(md):
        if kind == "heading":
            return _clean_inline(payload[1])
    return ""


# ---------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------

def _need(module: str, package: str):
    try:
        return __import__(module)
    except ImportError as exc:
        raise DocumentError(
            "Falta la libreria %s. Instalala con:  py -m pip install %s"
            % (package, package)) from exc


def _add_runs(paragraph, text: str, base_bold: bool = False) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and part.endswith(("**", "__")) \
                and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.bold = base_bold or None
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(part)
            if base_bold:
                run.bold = True


def markdown_to_docx(md: str, path, title: str = "") -> Path:
    _need("docx", "python-docx")
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.author = "Great Sage"
    if title:
        doc.core_properties.title = title
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    saw_title = False
    for kind, payload in _blocks(md):
        if kind == "heading":
            level, text = payload
            text = _clean_inline(text)
            if level == 1 and not saw_title:
                doc.add_heading(text, 0)          # the document's Title
                saw_title = True
            else:
                # "##" is a first-level section once a Title exists.
                lvl = level - 1 if saw_title or level > 1 else 1
                doc.add_heading(text, max(1, min(lvl, 4)))
        elif kind == "para":
            _add_runs(doc.add_paragraph(), payload)
        elif kind in ("bullet", "numbered"):
            depth, text = payload
            base = "List Bullet" if kind == "bullet" else "List Number"
            style = base if depth == 0 else "%s %d" % (base, min(depth + 1, 3))
            _add_runs(doc.add_paragraph(style=style), text)
        elif kind == "quote":
            _add_runs(doc.add_paragraph(style="Quote"), payload)
        elif kind == "code":
            for line in payload or [""]:
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                p.paragraph_format.space_after = Pt(0)
        elif kind == "table":
            _docx_table(doc, payload)
        elif kind == "rule":
            doc.add_paragraph()

    path = Path(path)
    doc.save(str(path))
    return path


def _docx_table(doc, rows: List[List[str]]) -> None:
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            text = row[c] if c < len(row) else ""
            _add_runs(cell.paragraphs[0], text, base_bold=(r == 0))
    doc.add_paragraph()


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------

_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")
_NUMBER_COMMAS = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")
_PERCENT = re.compile(r"^-?\d+(\.\d+)?\s*%$")
_MONEY = re.compile(r"^([$\u20ac\u00a3])\s*(-?\d[\d,]*(\.\d+)?)$")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _cell_value(raw: str):
    """(value, number_format). Formulas stay strings starting with '='."""
    text = _clean_inline(raw).strip()
    if text.startswith("="):
        return text, None
    if _NUMBER.match(text):
        return (int(text) if "." not in text else float(text)), None
    if _NUMBER_COMMAS.match(text):
        v = text.replace(",", "")
        return (int(v) if "." not in v else float(v)), "#,##0.##"
    if _PERCENT.match(text):
        return float(text.replace("%", "").strip()) / 100.0, "0.##%"
    m = _MONEY.match(text)
    if m:
        v = m.group(2).replace(",", "")
        return float(v), '"%s"#,##0.00' % m.group(1)
    return text, None


def _sheet_title(name: str, used: set) -> str:
    base = _SHEET_BAD.sub(" ", _clean_inline(name or "")).strip() or "Hoja"
    base = base[:31]
    title, n = base, 2
    while title.lower() in used:
        suffix = " %d" % n
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def markdown_to_xlsx(md: str, path, title: str = "") -> Path:
    _need("openpyxl", "openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets: List[Tuple[str, List[List[str]]]] = []
    heading = ""
    loose: List[str] = []
    for kind, payload in _blocks(md):
        if kind == "heading":
            heading = payload[1]
        elif kind == "table":
            sheets.append((heading, payload))
            heading = ""
        elif kind in ("para", "bullet", "numbered", "quote"):
            loose.append(payload if kind in ("para", "quote") else payload[1])

    if not sheets:
        # No table at all: keep the content, one line per row, rather than
        # handing back an empty workbook.
        rows = [[_clean_inline(t)] for t in loose] or [["(vacio)"]]
        sheets.append(("Datos", rows))

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    header_fill = PatternFill("solid", start_color="DDE6F3")

    for index, (name, rows) in enumerate(sheets, 1):
        ws = wb.create_sheet(_sheet_title(name or "Hoja %d" % index, used))
        widths: dict = {}
        for r, row in enumerate(rows, 1):
            for c, raw in enumerate(row, 1):
                value, fmt = _cell_value(raw)
                cell = ws.cell(row=r, column=c, value=value)
                if fmt:
                    cell.number_format = fmt
                if r == 1 and len(rows) > 1:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center")
                shown = len(str(value)) if not str(value).startswith("=") else 10
                widths[c] = max(widths.get(c, 0), shown)
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 3, 10), 60)
        if len(rows) > 1:
            ws.freeze_panes = "A2"

    wb.properties.creator = "Great Sage"
    if title:
        wb.properties.title = title
    path = Path(path)
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------

def _slides_from_markdown(md: str):
    """(deck_title, subtitle, [ {title, items:[(depth,text)], notes} ])."""
    deck_title, subtitle = "", ""
    slides: List[dict] = []
    current = None
    for kind, payload in _blocks(md):
        if kind == "heading":
            level, text = payload
            text = _clean_inline(text)
            if level == 1 and not deck_title and not slides:
                deck_title = text
            else:
                current = {"title": text, "items": [], "notes": ""}
                slides.append(current)
        elif current is None:
            # Text before the first slide is the deck's subtitle.
            if kind == "para" and not subtitle:
                subtitle = _clean_inline(payload)
        elif kind in ("bullet", "numbered"):
            current["items"].append((min(payload[0], 2), _clean_inline(payload[1])))
        elif kind == "para":
            text = _clean_inline(payload)
            m = re.match(r"^(?:notas?|notes?)\s*:\s*(.*)$", text, re.I)
            if m:
                current["notes"] = m.group(1)
            else:
                current["items"].append((0, text))
        elif kind == "table":
            for row in payload:
                current["items"].append((0, "  |  ".join(_clean_inline(c) for c in row)))
        elif kind == "quote":
            current["items"].append((0, _clean_inline(payload)))
    if not deck_title and slides:
        deck_title = slides[0]["title"]
    return deck_title, subtitle, slides


def markdown_to_pptx(md: str, path, title: str = "") -> Path:
    _need("pptx", "python-pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    NAVY = RGBColor(0x1F, 0x38, 0x64)
    ACCENT = RGBColor(0x2E, 0x75, 0xB6)
    INK = RGBColor(0x26, 0x26, 0x26)

    deck_title, subtitle, slides = _slides_from_markdown(md)
    if not slides and not deck_title:
        raise DocumentError("No encontre diapositivas (## Titulo) en el contenido.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height

    # -- title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    t, sub = s.shapes.title, s.placeholders[1]
    t.left, t.top, t.width, t.height = Inches(0.9), Inches(2.3), W - Inches(1.8), Inches(1.8)
    t.text_frame.text = deck_title or title or "Presentacion"
    t.text_frame.word_wrap = True
    for p in t.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size, r.font.bold = Pt(44), True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(4.3), W - Inches(1.8), Inches(1.2)
    sub.text_frame.text = subtitle
    for p in sub.text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(22)
            r.font.color.rgb = RGBColor(0xC9, 0xD6, 0xEA)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.15),
                             Inches(1.4), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    # -- content slides
    for spec in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        t, body = s.shapes.title, s.placeholders[1]
        t.left, t.top, t.width, t.height = Inches(0.8), Inches(0.45), W - Inches(1.6), Inches(1.1)
        t.text_frame.text = spec["title"]
        t.text_frame.word_wrap = True
        for p in t.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for r in p.runs:
                r.font.size, r.font.bold = Pt(34), True
                r.font.color.rgb = NAVY
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.55),
                                 Inches(1.2), Inches(0.07))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()

        body.left, body.top = Inches(0.8), Inches(1.95)
        body.width, body.height = W - Inches(1.6), H - Inches(2.6)
        tf = body.text_frame
        tf.word_wrap = True
        items = spec["items"] or [(0, "")]
        longest = max(len(x[1]) for x in items)
        size = 24 if (len(items) <= 5 and longest <= 80) else \
               20 if (len(items) <= 7 and longest <= 110) else 16
        for idx, (depth, text) in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = text
            p.level = depth
            p.space_after = Pt(10)
            for r in p.runs:
                r.font.size = Pt(size - 3 * depth)
                r.font.color.rgb = INK
        if spec["notes"]:
            s.notes_slide.notes_text_frame.text = spec["notes"]

    prs.core_properties.author = "Great Sage"
    prs.core_properties.title = deck_title or title
    path = Path(path)
    prs.save(str(path))
    return path


# ---------------------------------------------------------------------

_BUILDERS = {"docx": markdown_to_docx, "xlsx": markdown_to_xlsx,
             "pptx": markdown_to_pptx}


def build(kind: str, md: str, title: str = "",
          directory: Optional[Path] = None) -> Path:
    """Build the file for `kind` from Markdown and return its new path."""
    if kind not in _BUILDERS:
        raise DocumentError("Tipo de documento desconocido: %r" % kind)
    if not (md or "").strip():
        raise DocumentError("El contenido esta vacio; no hay nada que guardar.")
    heading = first_heading(md)
    name = title or heading or "documento"
    target = save_path(kind, name, directory)
    try:
        return _BUILDERS[kind](md, target, heading or title)
    except DocumentError:
        raise
    except Exception as exc:
        # Never leave a half-written file behind.
        try:
            if target.exists():
                os.remove(target)
        except OSError:
            pass
        raise DocumentError("No pude crear el archivo %s: %s" % (kind, exc)) from exc
